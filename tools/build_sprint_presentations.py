#!/usr/bin/env python3
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE

ROOT = Path(__file__).resolve().parent.parent

W, H = Inches(13.333), Inches(7.5)
NAVY = "0B1736"
BLUE = "2563EB"
CYAN = "0891B2"
TEAL = "0F766E"
ORANGE = "EA580C"
GREEN = "15803D"
RED = "B91C1C"
PURPLE = "7C3AED"
INK = "111827"
MUTED = "64748B"
LINE = "DCE3EC"
PAPER = "F7F9FC"
WHITE = "FFFFFF"
PALE_BLUE = "EAF2FF"
PALE_CYAN = "E6F7FA"
PALE_GREEN = "EAF7EE"
PALE_ORANGE = "FFF1E8"
PALE_RED = "FDECEC"
PALE_PURPLE = "F3EEFF"
FONT = "TH Sarabun New"

def rgb(hexstr):
    return RGBColor.from_string(hexstr)

def set_bg(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)

def add_text(slide, text, x, y, w, h, size=20, bold=False, color=INK,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box

def add_bullets(slide, items, x, y, w, h, size=18, color=INK, bullet_color=None,
                gap=7, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = font
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(gap)
        p.line_spacing = 1.1
        p.text = "• " + item
    return box

def rect(slide, x, y, w, h, fill, line=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    return shape

def line(slide, x1, y1, x2, y2, color=LINE, width=1.5, dash=None):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return ln

def metric_card(slide, x, y, w, value, label, fill=WHITE, accent=BLUE):
    rect(slide, x, y, w, 1.18, fill, LINE)
    rect(slide, x, y, 0.10, 1.18, accent, accent, radius=False)
    add_text(slide, value, x+0.25, y+0.12, w-0.35, 0.52, 26, True, accent)
    add_text(slide, label, x+0.25, y+0.66, w-0.35, 0.34, 12, False, MUTED)
    
def header(slide, sprint, title, subtitle=None, accent=BLUE):
    add_text(slide, sprint.upper(), 0.58, 0.30, 1.7, 0.35, 11, True, accent)
    add_text(slide, title, 0.58, 0.68, 11.9, 0.55, 25, True, NAVY)
    if subtitle:
        add_text(slide, subtitle, 0.58, 1.20, 11.9, 0.34, 12.5, False, MUTED)
    line(slide, 0.58, 1.63, 12.75, 1.63, LINE, 1.2)

def footer(slide, source, num, accent=BLUE):
    line(slide, 0.58, 7.06, 12.75, 7.06, LINE, 0.8)
    add_text(slide, source, 0.58, 7.08, 10.8, 0.24, 9, False, MUTED)
    add_text(slide, f"{num:02d}", 12.0, 7.06, 0.7, 0.24, 10, True, accent, PP_ALIGN.RIGHT)

def title_slide(prs, d):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, NAVY)
    rect(slide, 0, 0, 13.333, 0.11, d["accent"], d["accent"], radius=False)
    add_text(slide, f"SPRINT {d['num']}", 0.72, 0.65, 2.0, 0.36, 13, True, d["accent"])
    add_text(slide, d["title"], 0.72, 1.25, 11.8, 1.4, 31, True, WHITE, valign=MSO_ANCHOR.TOP)
    add_text(slide, d["subtitle"], 0.75, 2.75, 11.5, 0.62, 17, False, "C7D2FE", valign=MSO_ANCHOR.TOP)
    for i, (v,l) in enumerate(d["hero_metrics"]):
        x = 0.75 + i*3.95
        rect(slide, x, 4.25, 3.55, 1.25, "13244B", "28406F")
        add_text(slide, v, x+0.22, 4.43, 3.0, 0.42, 27, True, WHITE)
        add_text(slide, l, x+0.22, 4.92, 3.0, 0.34, 12, False, "B8C4E4")
    add_text(slide, d["period"], 0.75, 6.55, 5.4, 0.32, 11.5, False, "B8C4E4")
    add_text(slide, "Software Project Management · ข้อมูลสมมติเพื่อการเรียน", 6.25, 6.55, 6.3, 0.32, 10.5, False, "8393B6", PP_ALIGN.RIGHT)
    return slide

def overview_slide(prs, d):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
    header(slide, f"Sprint {d['num']}", "ภาพรวมสปรินต์", d["goal"], d["accent"])
    metrics = d["overview_metrics"]
    for i, m in enumerate(metrics):
        metric_card(slide, 0.62+i*3.15, 1.95, 2.86, *m, accent=d["accent"])
    add_text(slide, "ทีมและความรับผิดชอบ", 0.65, 3.45, 5.4, 0.38, 17, True, NAVY)
    y = 3.92
    for name, role, task in d["team"]:
        rect(slide, 0.65, y, 5.82, 0.72, WHITE, LINE)
        add_text(slide, name, 0.86, y+0.08, 2.1, 0.26, 12.5, True, INK)
        add_text(slide, role, 0.86, y+0.34, 2.1, 0.22, 10.5, False, d["accent"])
        add_text(slide, task, 3.05, y+0.08, 3.15, 0.48, 10.7, False, MUTED, valign=MSO_ANCHOR.TOP)
        y += 0.86
    add_text(slide, "เป้าหมายหลัก", 6.85, 3.45, 5.4, 0.38, 17, True, NAVY)
    add_bullets(slide, d["goal_bullets"], 6.85, 3.88, 5.6, 2.45, 15.5, INK, gap=11)
    footer(slide, d["source"], 2, d["accent"])

def timeline_slide(prs, d):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
    header(slide, f"Sprint {d['num']}", d["timeline_title"], d["timeline_subtitle"], d["accent"])
    items = d["timeline"]
    x0, x1 = 1.05, 12.25
    y = 3.2
    line(slide, x0, y, x1, y, "B9C6D8", 3)
    n = len(items)
    for i, (tag, title, desc) in enumerate(items):
        x = x0 + (x1-x0) * (i/(n-1 if n>1 else 1))
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-0.16), Inches(y-0.16), Inches(0.32), Inches(0.32))
        circ.fill.solid(); circ.fill.fore_color.rgb = rgb(d["accent"]); circ.line.color.rgb = rgb(d["accent"])
        add_text(slide, tag, x-0.65, 2.35, 1.3, 0.38, 13, True, d["accent"], PP_ALIGN.CENTER)
        add_text(slide, title, x-1.25, 3.52, 2.5, 0.44, 14, True, NAVY, PP_ALIGN.CENTER)
        add_text(slide, desc, x-1.35, 4.02, 2.7, 1.28, 11.1, False, MUTED, PP_ALIGN.CENTER, MSO_ANCHOR.TOP)
    footer(slide, d["source"], 3, d["accent"])

def cards_slide(prs, d, title, subtitle, cards, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
    header(slide, f"Sprint {d['num']}", title, subtitle, d["accent"])
    cols = 2
    card_w, card_h = 5.86, 1.82
    for idx, c in enumerate(cards):
        col, row = idx%cols, idx//cols
        x = 0.65 + col*6.05
        y = 1.95 + row*2.08
        fill = c.get("fill", WHITE)
        rect(slide, x, y, card_w, card_h, fill, LINE)
        rect(slide, x, y, 0.10, card_h, c.get("accent", d["accent"]), c.get("accent", d["accent"]), radius=False)
        add_text(slide, c["title"], x+0.28, y+0.16, 5.2, 0.38, 16, True, NAVY)
        if c.get("metric"):
            add_text(slide, c["metric"], x+4.2, y+0.16, 1.25, 0.38, 14, True, c.get("accent", d["accent"]), PP_ALIGN.RIGHT)
        add_text(slide, c["body"], x+0.28, y+0.60, 5.18, 1.02, 11.4, False, MUTED, valign=MSO_ANCHOR.TOP)
    footer(slide, d["source"], num, d["accent"])

def compare_slide(prs, d, title, left_title, left_items, right_title, right_items, num, center=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
    header(slide, f"Sprint {d['num']}", title, None, d["accent"])
    rect(slide, 0.68, 2.0, 5.65, 4.35, WHITE, LINE)
    rect(slide, 7.0, 2.0, 5.65, 4.35, WHITE, LINE)
    add_text(slide, left_title, 0.95, 2.22, 5.0, 0.4, 19, True, RED)
    add_text(slide, right_title, 7.28, 2.22, 5.0, 0.4, 19, True, GREEN)
    add_bullets(slide, left_items, 0.98, 2.82, 4.95, 2.9, 15, INK, gap=12)
    add_bullets(slide, right_items, 7.3, 2.82, 4.95, 2.9, 15, INK, gap=12)
    if center:
        rect(slide, 6.22, 3.55, 0.88, 0.72, d["accent"], d["accent"])
        add_text(slide, center, 6.24, 3.63, 0.84, 0.50, 16, True, WHITE, PP_ALIGN.CENTER)
    footer(slide, d["source"], num, d["accent"])

def metrics_slide(prs, d, title, subtitle, metrics, notes, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
    header(slide, f"Sprint {d['num']}", title, subtitle, d["accent"])
    for i, (value,label,accent) in enumerate(metrics):
        metric_card(slide, 0.7+i*3.15, 2.0, 2.85, value, label, fill=WHITE, accent=accent)
    rect(slide, 0.7, 3.52, 12.0, 2.70, WHITE, LINE)
    add_text(slide, "อ่านค่าร่วมกัน", 0.98, 3.78, 3.0, 0.38, 16.5, True, NAVY)
    add_bullets(slide, notes, 0.98, 4.22, 11.2, 1.65, 14.6, INK, gap=11)
    footer(slide, d["source"], num, d["accent"])

def review_slide(prs, d, done, risks, nexts, num=8):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
    header(slide, f"Sprint {d['num']}", "Sprint Review & Handoff", "สิ่งที่ยืนยันได้จากหลักฐาน และสิ่งที่ส่งต่อไปยังรอบถัดไป", d["accent"])
    blocks = [
        ("ส่งมอบแล้ว", done, PALE_GREEN, GREEN),
        ("ความเสี่ยง/ข้อจำกัด", risks, PALE_RED, RED),
        ("งานส่งต่อ", nexts, PALE_BLUE, d["accent"]),
    ]
    for i,(title,items,fill,accent) in enumerate(blocks):
        x = 0.65 + i*4.05
        rect(slide, x, 1.95, 3.82, 4.7, fill, LINE)
        add_text(slide, title, x+0.26, 2.2, 3.2, 0.4, 18, True, accent)
        add_bullets(slide, items, x+0.26, 2.78, 3.25, 3.25, 13.1, INK, gap=11)
    footer(slide, d["source"], num, d["accent"])

def architecture_slide(prs, d, title, subtitle, lanes, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
    header(slide, f"Sprint {d['num']}", title, subtitle, d["accent"])
    y = 2.05
    widths = [2.15, 2.55, 2.55, 2.35]
    x = 0.75
    centers=[]
    for i, lane in enumerate(lanes):
        w=widths[i] if i < len(widths) else 2.2
        rect(slide, x, y, w, 2.05, lane.get("fill", WHITE), LINE)
        add_text(slide, lane["title"], x+0.18, y+0.25, w-0.36, 0.42, 16, True, lane.get("accent", d["accent"]), PP_ALIGN.CENTER)
        add_text(slide, lane["body"], x+0.22, y+0.82, w-0.44, 0.92, 11.3, False, MUTED, PP_ALIGN.CENTER, MSO_ANCHOR.TOP)
        centers.append((x+w/2,y+1.02))
        if i < len(lanes)-1:
            add_text(slide, "→", x+w+0.05, y+0.74, 0.42, 0.5, 24, True, "94A3B8", PP_ALIGN.CENTER)
        x += w+0.55
    if d.get("architecture_notes"):
        rect(slide, 0.78, 4.65, 11.9, 1.45, PALE_BLUE, "D9E6FF")
        add_bullets(slide, d["architecture_notes"], 1.02, 4.86, 11.3, 0.98, 13.7, INK, gap=7)
    footer(slide, d["source"], num, d["accent"])

def build_sprint(d):
    prs = Presentation()
    prs.slide_width = W; prs.slide_height = H
    title_slide(prs, d)
    overview_slide(prs, d)
    timeline_slide(prs, d)
    # slide 4
    s4=d["slide4"]
    if s4["type"]=="compare":
        compare_slide(prs,d,s4["title"],s4["left_title"],s4["left_items"],s4["right_title"],s4["right_items"],4,s4.get("center"))
    elif s4["type"]=="architecture":
        d["architecture_notes"]=s4.get("notes",[])
        architecture_slide(prs,d,s4["title"],s4.get("subtitle"),s4["lanes"],4)
    else:
        cards_slide(prs,d,s4["title"],s4.get("subtitle"),s4["cards"],4)
    # slide 5
    s5=d["slide5"]
    if s5["type"]=="cards":
        cards_slide(prs,d,s5["title"],s5.get("subtitle"),s5["cards"],5)
    elif s5["type"]=="metrics":
        metrics_slide(prs,d,s5["title"],s5.get("subtitle"),s5["metrics"],s5["notes"],5)
    else:
        compare_slide(prs,d,s5["title"],s5["left_title"],s5["left_items"],s5["right_title"],s5["right_items"],5,s5.get("center"))
    # slide 6
    s6=d["slide6"]
    if s6["type"]=="metrics":
        metrics_slide(prs,d,s6["title"],s6.get("subtitle"),s6["metrics"],s6["notes"],6)
    else:
        cards_slide(prs,d,s6["title"],s6.get("subtitle"),s6["cards"],6)
    # slide 7
    s7=d["slide7"]
    cards_slide(prs,d,s7["title"],s7.get("subtitle"),s7["cards"],7)
    review_slide(prs,d,d["done"],d["risks"],d["nexts"],8)
    out=ROOT/d["out"]
    out.parent.mkdir(parents=True,exist_ok=True)
    prs.save(out)
    print("pptx", out.relative_to(ROOT))
    return out

TEAM = [
    ("ภานุวัฒน์ ต๋าคำ","PM / Developer","บริหารภาพรวม พัฒนา และจัดการหลักฐานโครงการ"),
    ("เอกพันธ์ ทศทิศรังสรรค์","QA / Tester","ออกแบบและยืนยันผลทดสอบ"),
    ("ณฐภาพ สายหล้า","Tech Lead / Architect","ออกแบบสถาปัตยกรรมและตรวจคุณภาพ"),
]

SPRINTS = [
{
"num":1,"accent":BLUE,
"title":"เริ่มโครงการและวางฐาน",
"subtitle":"จากระบบ Monolithic → โครงสร้าง OOP ที่ทดสอบได้และปลอดภัยขึ้น",
"period":"Phase 1 · W1–W4 · Jira Sprint ID 84",
"hero_metrics":[("15 SP","Story Points ปิดครบ"),("5 Issues","SPM-6 ถึง SPM-10"),("5/5","PyTest ผ่าน")],
"overview_metrics":[("W1–W4","ช่วงงาน",WHITE),("15 SP","ปิดครบ",WHITE),("5","Issues",WHITE),("5/5","Tests PASS",WHITE)],
"team":TEAM,
"goal":"ตั้งกฎบัตร เข้าใจระบบเดิม ออกแบบพิมพ์เขียว และ refactor พร้อม automated test",
"goal_bullets":["กำหนด Project Charter / Scope และความเสี่ยง","วิเคราะห์ DFD, hotspot และ static quality ของระบบเดิม","ออกแบบ blueprint สำหรับ OOP + Member/Discount","ลงมือ refactor, validation, atomic save และ PyTest"],
"timeline_title":"เส้นทาง W1–W4",
"timeline_subtitle":"งานจากการตั้งต้นโครงการไปจนถึงโค้ด refactor ที่ทดสอบซ้ำได้",
"timeline":[
("W1","ตั้งต้น","Charter · Scope · ระบบเดิม 5 เมนู · ความเสี่ยง"),
("W2","วิเคราะห์","DFD · Hotspot · Static Analysis · Blueprint"),
("W3","พัฒนา","OOP 3 คลาส · Validation · Atomic Save · DoD/RACI"),
("W4","ยืนยันผล","ปรับโค้ดรอบสุดท้าย · PyTest 5 เคสผ่าน")
],
"slide4":{"type":"compare","title":"ก่อนและหลังการ Refactor",
"left_title":"As-Is: v1","left_items":["Monolithic + global state","JSON เขียนทับตรง เสี่ยงข้อมูลเสีย","ไม่มี validation ค่าติดลบ/ชนิดข้อมูล","main() ซับซ้อน: CC = 14"],
"right_title":"Sprint 1 Deliverable","right_items":["Product / InventoryManager / InventoryCLI","Atomic write: temp + os.replace","Validator แยกตามช่อง + กัน ValueError","Business Logic แยกจาก CLI"],"center":"→"},
"slide5":{"type":"metrics","title":"Baseline คุณภาพโค้ด","subtitle":"ค่าก่อนปรับใช้เป็นฐานเปรียบเทียบ — ไม่อ้างเป็นค่าหลังปรับ",
"metrics":[("99","Lines of Code",BLUE),("7.10","Pylint / 10",CYAN),("14","Cyclomatic Complexity",ORANGE),("17","Branches",RED)],
"notes":["main() มี statements 55 (>50) และ branch 17 (>12)","หลัง refactor โครงสร้างดีขึ้นและเทสต์ผ่าน แต่ Sprint 1 ไม่ได้วัด Pylint/CC ซ้ำ","เกณฑ์เป้าหมายในแผนถัดไป: ลดความซับซ้อนและแยกหน้าที่ให้ชัด"]},
"slide6":{"type":"metrics","title":"ผลทดสอบและ EVM","subtitle":"ยืนยันทั้งด้านเทคนิคและภาพการบริหารรอบแรก",
"metrics":[("5/5","PyTest PASS",GREEN),("4,000","PV (THB)",BLUE),("3,000","EV (THB)",CYAN),("4,200","AC (THB)",ORANGE)],
"notes":["SV = −1,000 THB → งานช้ากว่าแผน","CV = −1,200 THB → ใช้ต้นทุนมากกว่าคุณค่าที่ได้","สาเหตุหลัก: technical debt ซ่อนใน app_v1.py และเวลาแก้ merge conflict"]},
"slide7":{"title":"สิ่งที่เรียนรู้จาก Sprint 1","subtitle":"ฐานที่ใช้กำหนดวิธีทำงานใน Sprint ถัดไป",
"cards":[
{"title":"คุณภาพ","body":"DoD 5 ข้อ + RACI ทำให้การส่งมอบมีเกณฑ์ตรวจที่ชัดเจน","accent":GREEN,"fill":PALE_GREEN},
{"title":"ความเสี่ยงข้อมูล","body":"Atomic write ลดความเสี่ยง JSON พัง แต่ยังไม่มี transaction ระดับฐานข้อมูล","accent":ORANGE,"fill":PALE_ORANGE},
{"title":"Backward Compatibility","body":"รองรับคีย์เก่า n/q/p/c และค่า default เมื่อฟิลด์ขาด","accent":CYAN,"fill":PALE_CYAN},
{"title":"งานยังไม่จบ","body":"SQLite, Member tiers และ Checkout ยังเป็นแบบออกแบบ ต้องกลับมาทำจริง","accent":PURPLE,"fill":PALE_PURPLE},
]},
"done":["OOP 3 คลาสและ separation of concerns","Atomic save + input validation","PyTest 5 เคสพร้อม test.json"],
"risks":["JSON ยังไม่มี transaction","E501 ค้างใน snapshot ประวัติ","ไม่มี formal sponsor sign-off"],
"nexts":["พัฒนา SQLite + Member + Checkout","รับ CR-01 / CR-02 และแก้ BUG-101","เตรียม UAT และ release gate"],
"source":"Source: Phase1/Sprint1/sprint1.md",
"out":"Phase1/Sprint1/sprint1-presentation.pptx"
},
{
"num":2,"accent":CYAN,
"title":"ออกแบบและตั้งต้นทุนฐาน",
"subtitle":"วาง To-Be Architecture, Quality Gate และ Cost Baseline ก่อนลงมือจริง",
"period":"Phase 2 · W5–W7 · Jira Sprint ID 82",
"hero_metrics":[("13 SP","Story Points ปิดครบ"),("18,823 THB","Cost Baseline"),("≥ 85%","Coverage Target")],
"overview_metrics":[("W5–W7","ช่วงงาน",WHITE),("13 SP","ปิดครบ",WHITE),("4","Issues",WHITE),("36 h","Capacity / Sprint",WHITE)],
"team":[
("ภานุวัฒน์ ต๋าคำ","PM / Developer","ประมาณการต้นทุน · budget worksheet · WBS"),
("เอกพันธ์ ทศทิศรังสรรค์","QA / Tester","กำหนด coverage และ quality gate"),
("ณฐภาพ สายหล้า","Tech Lead / Architect","ออกแบบ To-Be และมาตรการ CI"),
],
"goal":"ออกแบบสถาปัตยกรรมเป้าหมายและกำหนดกรอบคุณภาพ/งบประมาณที่ใช้ควบคุมการพัฒนาจริง",
"goal_bullets":["เทียบ As-Is กับ To-Be และประเมิน ISO 25010","ออกแบบ SQLite + Repository + Strategy + Member tier","กำหนด Communication Matrix, DoR และ DoD","อนุมัติ Cost Baseline 18,823 THB + Threshold คุมงบ"],
"timeline_title":"เส้นทาง W5–W7",
"timeline_subtitle":"เปลี่ยนจากแนวคิดสถาปัตยกรรมให้เป็นแผนที่วัดและควบคุมได้",
"timeline":[
("W5","ประเมิน","As-Is / To-Be · ISO 25010 · Budget worksheet"),
("W6","ออกแบบ","Communication Matrix · DoR/DoD · Architecture · Variance"),
("W7","Baseline","Cost Baseline · Capacity · ISO 14598 Quality Gate"),
],
"slide4":{"type":"architecture","title":"To-Be Architecture ที่ออกแบบไว้","subtitle":"แยกความรับผิดชอบและเตรียมทางย้ายจาก JSON ไป SQLite",
"lanes":[
{"title":"Presentation","body":"CLI / UI\nรับข้อมูลและแสดงผล","accent":BLUE,"fill":PALE_BLUE},
{"title":"Service","body":"Checkout\nValidation\nExport","accent":CYAN,"fill":PALE_CYAN},
{"title":"Repository","body":"ซ่อน SQL\nCRUD / Summary\nParameterized Query","accent":PURPLE,"fill":PALE_PURPLE},
{"title":"SQLite","body":"Products\nMembers\nTransaction","accent":GREEN,"fill":PALE_GREEN},
],
"notes":["Singleton จัดการ connection · Repository ซ่อน persistence · Strategy รองรับ Member tier","แบบออกแบบนี้ถูกนำไปพัฒนาจริงใน Sprint 5"]},
"slide5":{"type":"metrics","title":"ISO/IEC 25010: As-Is","subtitle":"4 มิติผ่าน · 1 มิติต้องปรับ · 3 มิติตกและต้องแก้เชิงโครงสร้าง",
"metrics":[("4","ผ่าน",GREEN),("1","ต้องปรับ",ORANGE),("3","ตก",RED),("8","มิติทั้งหมด",BLUE)],
"notes":["Reliability: JSON เขียนทับตรง → เป้าหมาย atomic write/SQLite","Security: เตรียม parameterized query 100% เมื่อเข้าสู่ SQL","Maintainability: Monolithic / CC 14 → เป้าหมายแยกชั้น + OOP","Usability: เพิ่ม validator และลดความซับซ้อนของเมนู"]},
"slide6":{"type":"metrics","title":"Cost Baseline และ Quality Gate","subtitle":"ตัวเลขที่อนุมัติใน W7 ใช้เป็นฐานควบคุมงบ",
"metrics":[("15,000","Labor (THB)",BLUE),("1,000","Infra (THB)",CYAN),("2,823","Reserve (THB)",ORANGE),("18,823","Approved Total",GREEN)],
"notes":["Threshold: ≤5% ปกติ · 5–10% แจ้งเตือน · >10% วิกฤต/พิจารณาตัด scope","Capacity: 3 คน × 6 ชั่วโมง × 2 สัปดาห์ = 36 man-hours","Quality Gate: v(G) ≤ 8 · CBO ≤ 4 · DIT ≤ 2 · Coverage ≥ 85%"]},
"slide7":{"title":"ความเสี่ยงของแผน","subtitle":"จุดที่ต้องระวังเมื่อนำ baseline ไปใช้จริง",
"cards":[
{"title":"ตัวเลขต่างฐาน","body":"W5 มี worksheet 39,100 THB แต่ W7 อนุมัติ 18,823 THB — ต้องระบุฐานทุกครั้ง","accent":RED,"fill":PALE_RED},
{"title":"Capacity","body":"36 man-hours ต่อ Sprint ต่ำกว่างาน evolution จริงที่พบในภายหลัง","accent":ORANGE,"fill":PALE_ORANGE},
{"title":"Quality Target","body":"ISO 14598 ยังเป็นเป้าหมาย ต้องวัดจริงช่วง hardening","accent":PURPLE,"fill":PALE_PURPLE},
{"title":"Cost Variance","body":"ตัวอย่าง migration: PV/EV 2,400 แต่ AC 4,000 → CV −1,600","accent":CYAN,"fill":PALE_CYAN},
]},
"done":["To-Be Architecture + DB design","Cost Baseline / Threshold / Capacity","Communication Matrix + DoR/DoD + Quality Gate"],
"risks":["งบ W5/W7 คนละฐาน","Capacity ต่ำกว่างานจริง","Quality target ยังไม่ได้วัดจริง"],
"nexts":["รับ CR-01 ผ่าน Impact Analysis","ติดตามด้วย Burndown / EVM","วัด ISO gate ในรอบ Hardening"],
"source":"Source: Phase2/Sprint2/sprint2.md",
"out":"Phase2/Sprint2/sprint2-presentation.pptx"
},
{
"num":3,"accent":ORANGE,
"title":"ลงมือและคุมการเปลี่ยนแปลง",
"subtitle":"Change Control · EVM · Defect-driven Test · Hardening · Scope Freeze",
"period":"Phase 3 · W8–W11 · Jira Sprint ID 85",
"hero_metrics":[("16 SP","Story Points ปิดครบ"),("27 / 29","Burnup"),("1,350 THB","Reserve Used")],
"overview_metrics":[("W8–W11","ช่วงงาน",WHITE),("16 SP","ปิดครบ",WHITE),("5","Issues",WHITE),("93%","Burnup",WHITE)],
"team":[
("ภานุวัฒน์ ต๋าคำ","PM / Developer","Impact Analysis · CCB · Log Work · EVM"),
("เอกพันธ์ ทศทิศรังสรรค์","QA / Tester","Defect-driven / integration test"),
("ณฐภาพ สายหล้า","Tech Lead / Architect","Root cause · review · hardening"),
],
"goal":"นำแผนไปใช้จริง พร้อมควบคุม Change Request, defect, งบ และขอบเขตไม่ให้หลุด",
"goal_bullets":["รับ CR-01 ผ่าน Impact Analysis และ board tracking","ติดตาม EVM + retrospective + transition","อนุมัติ CR-02 ผ่าน CCB และแก้ BUG-101 ด้วย 5 Whys","Hardening โค้ดและ Scope Freeze ก่อนเข้าสู่ release"],
"timeline_title":"เส้นทาง W8–W11",
"timeline_subtitle":"ทุกการเปลี่ยนแปลงต้องมีผลกระทบ หลักฐาน และเงื่อนไขปิด",
"timeline":[
("W8","CR-01","Impact · Burndown · Blocker · Stakeholder"),
("W9","ติดตาม","EVM · Retro · Sprint Transition"),
("W10","CR-02 + BUG","CCB · 5 Whys · Reserve"),
("W11","Harden","Flow/WIP · Burnup · Scope Freeze"),
],
"slide4":{"type":"compare","title":"Change Request ที่ควบคุมผ่านกระบวนการ",
"left_title":"CR-01 · Barcode + ROP","left_items":["Perfective / Normal Change","+8 man-hours impact","แตะ Product + Repository + ConsoleUI","เพิ่ม test สำหรับ barcode / reorder point"],
"right_title":"CR-02 · CSV Export","right_items":["Emergency / Fast-Track","4.5 h × 300 = 1,350 THB","เพิ่ม CsvReportExporter ตาม SRP","อนุมัติผ่าน CCB และดึง reserve"],"center":"CCB"},
"slide5":{"type":"cards","title":"BUG-101: จาก Defect สู่ Permanent Test","subtitle":"แก้ที่ root cause ไม่ใช่แค่ปิดอาการ",
"cards":[
{"title":"อาการ","body":"KeyError: 'barcode' เมื่อโหลด data.json รุ่นเก่าที่ไม่มีฟิลด์ใหม่","accent":RED,"fill":PALE_RED},
{"title":"Root Cause","body":"Schema ใหม่ไม่มี Data Validation/default fallback ในชั้น Repository","accent":ORANGE,"fill":PALE_ORANGE},
{"title":"Fix","body":"dict.get('barcode','') และ reorder_point default = 5","accent":GREEN,"fill":PALE_GREEN},
{"title":"Prevention","body":"Defect-driven test + backward compatibility test เป็น regression ถาวร","accent":CYAN,"fill":PALE_CYAN},
]},
"slide6":{"type":"metrics","title":"Progress / Flow / Cost","subtitle":"ตัวเลขสำคัญก่อนแช่แข็งขอบเขต",
"metrics":[("27 / 29","Burnup",GREEN),("3","WIP Limit: Review",BLUE),("1,350","Reserve Used",ORANGE),("1,473","Reserve Left",CYAN)],
"notes":["Burnup = 93% เหลือ 2 points ก่อน release","WIP Limit 3 ใบช่วยลดงานค้างใน Review","EVM Sprint 1 ที่ติดตามย้อนหลัง: PV 4,000 · EV 3,000 · AC 4,200","Scope Freeze อนุญาตหลัง freeze เฉพาะ Critical bug / Security fix / Test"]},
"slide7":{"title":"Retrospective: Mad / Sad / Glad","subtitle":"บทเรียนถูกเปลี่ยนเป็น action ที่วัดได้",
"cards":[
{"title":"Mad","body":"เสียเวลา merge conflict จากการตัดโครงสร้างใหญ่ในครั้งเดียว","accent":RED,"fill":PALE_RED},
{"title":"Sad","body":"ประเมิน refactor ต่ำกว่าจริง เพราะ technical debt ใน app_v1.py","accent":ORANGE,"fill":PALE_ORANGE},
{"title":"Glad","body":"PyTest จับ defect ได้ก่อน demo และทีมแก้ conflict สำเร็จ","accent":GREEN,"fill":PALE_GREEN},
{"title":"Action","body":"แจ้ง blocker ทุกวัน · PR review ≤12h · ใช้ TDR/เขียนเทสต์ก่อนโค้ด","accent":BLUE,"fill":PALE_BLUE},
]},
"done":["CR-01 / CR-02 ทำงานจริง","BUG-101 แก้พร้อม regression test","Hardening + Scope Freeze 29 points"],
"risks":["เหลือ 2 points ใน burnup","E501 ยังอยู่ ณ W11","UAT formal sign-off ยังรอ"],
"nexts":["ปิด 29/29 และรัน UAT SC01–03","ทำ Release Checklist + Final EVM","แก้ E501 และสร้าง release tag"],
"source":"Source: Phase3/Sprint3/sprint3.md",
"out":"Phase3/Sprint3/sprint3-presentation.pptx"
},
{
"num":4,"accent":GREEN,
"title":"UAT และปล่อยรุ่น v2.0",
"subtitle":"ปิดวงจร W1–W12 ด้วย User Acceptance, Release Gate และ Final EVM",
"period":"Phase 4 · W12 · Jira Sprint ID 83",
"hero_metrics":[("5 SP","SPM-20 + SPM-21"),("25/25","Regression Tests"),("v2.0.0","Evolution Release")],
"overview_metrics":[("W12","ช่วงงาน",WHITE),("5 SP","ปิดครบ",WHITE),("2","Issues",WHITE),("25","Tests",WHITE)],
"team":[
("ภานุวัฒน์ ต๋าคำ","PM / Developer","Release checklist · Final EVM · procurement"),
("เอกพันธ์ ทศทิศรังสรรค์","QA / Tester","UAT scenarios · cross-team test"),
("ณฐภาพ สายหล้า","Tech Lead / Architect","Release gate · merge main"),
],
"goal":"ยืนยันว่าฟีเจอร์ evolution พร้อมใช้งานจริงทางเทคนิค และปิดสถานะโครงการด้าน release/งบ",
"goal_bullets":["รวม CR-01, CR-02 และ BUG-101 fix เข้ารุ่น v2.0","ขยาย test suite จาก 5 → 25 เคส","รัน UAT SC01–SC03 และตรวจ release gate 3 ด่าน","สรุป Final EVM และออก annotated tag"],
"timeline_title":"Release Flow ใน W12",
"timeline_subtitle":"จาก code complete → acceptance → technical gate → release",
"timeline":[
("Code","รวมฟีเจอร์","Barcode · ROP · CSV · BUG-101"),
("Test","Regression","25 เคสผ่าน"),
("UAT","SC01–03","เพิ่มสินค้า · ตัดสต็อก · Export CSV"),
("Release","v2.0.0","Merge · Tag · Release Note · Final EVM"),
],
"slide4":{"type":"compare","title":"วิวัฒนาการจาก v1.0 → v2.0",
"left_title":"v1.0","left_items":["JSON เขียนทับตรง","Monolithic + global x","ไม่มี Barcode / CSV","5 tests"],
"right_title":"v2.0.0-evolution","right_items":["JSON + atomic write","OOP 3 คลาส แยก Logic/UI","Barcode + Reorder Point + CSV","25 tests"],"center":"→"},
"slide5":{"type":"cards","title":"UAT SC01–SC03","subtitle":"ผลทางเทคนิคผ่านครบ แต่ formal sign-off ยังรอลายเซ็น",
"cards":[
{"title":"SC01 · Inventory Manager","metric":"PASS","body":"เพิ่ม Milk (885) จำนวน 10 · ROP 5 → บันทึกสำเร็จ ไม่ crash","accent":GREEN,"fill":PALE_GREEN},
{"title":"SC02 · Store Cashier","metric":"PASS","body":"ตัดสต็อก 6 เหลือ 4 → ระบบแจ้ง low stock ทันที","accent":GREEN,"fill":PALE_GREEN},
{"title":"SC03 · Purchasing Officer","metric":"PASS","body":"Export CSV เปิดใน Excel → มี Barcode และข้อมูลถูกต้อง","accent":GREEN,"fill":PALE_GREEN},
{"title":"Formal Sign-off","metric":"WAIT","body":"ช่องลายเซ็นผู้ทดสอบและ Tech Lead ยังว่าง จึงไม่อ้างว่ารับรองทางการเสร็จแล้ว","accent":ORANGE,"fill":PALE_ORANGE},
]},
"slide6":{"type":"metrics","title":"Final EVM และเงินสำรอง","subtitle":"ส่งงานตรงเวลา แต่ต้นทุนจริงสูงกว่าแผนเล็กน้อย",
"metrics":[("1.00","SPI",GREEN),("0.96","CPI",ORANGE),("−575","CV (THB)",RED),("898","Reserve Left",BLUE)],
"notes":["PV = EV = 15,525 THB → เสร็จตามแผน","AC = 16,100 THB → เกิน 575 THB หรือ 3.7%","Reserve: 2,823 − 1,350 (CR-02) − 575 (close-out) = 898 THB"]},
"slide7":{"title":"Release Checklist และ Changelog","subtitle":"สิ่งที่ถูกยืนยันก่อน/หลัง tag",
"cards":[
{"title":"Gate 1 · Quality","body":"UAT SC01–03 ผ่านและแยกประเด็น defect/change request","accent":GREEN,"fill":PALE_GREEN},
{"title":"Gate 2 · Technical","body":"Regression 25/25 · PR develop→main · CI · Tech Lead merge","accent":BLUE,"fill":PALE_BLUE},
{"title":"Gate 3 · Release","body":"Merge main · annotated tag · release note · Final EVM","accent":PURPLE,"fill":PALE_PURPLE},
{"title":"CHANGELOG","body":"Added Barcode/ROP/CSV · Changed schema/menu · Fixed BUG-101","accent":CYAN,"fill":PALE_CYAN},
]},
"done":["v2.0.0-evolution พร้อม tag","25 regression tests ผ่าน","UAT 3 scenarios ผ่านทางเทคนิค"],
"risks":["Formal UAT sign-off ยังว่าง","Legacy data ยังเป็น fallback","SQLite/Member ยังไม่ได้ implement"],
"nexts":["พัฒนา SQLite + Member + Checkout","แก้ E501 ในโค้ดสด","ต่อยอดเป็น Desktop UI โดย reuse domain"],
"source":"Source: Phase4/Sprint4/sprint4.md",
"out":"Phase4/Sprint4/sprint4-presentation.pptx"
},
{
"num":5,"accent":PURPLE,
"title":"วิวัฒนาการ v3.0 และโปรแกรม Desktop",
"subtitle":"SQLite + Member + Checkout → GTK4 Desktop UI โดยไม่ทำ regression",
"period":"Phase 5 · Jira Sprint ID 86 · 2 Epics",
"hero_metrics":[("46 SP","Story Points ปิดครบ"),("21 + 25","Current + Regression Tests"),("Desktop","python3 program.py")],
"overview_metrics":[("46 SP","ปิดครบ",WHITE),("9","Issues",WHITE),("2","Epics",WHITE),("0","Flake8/Bandit",WHITE)],
"team":[
("ภานุวัฒน์ ต๋าคำ","PM / Developer","Desktop UI · Migration · Checkout"),
("เอกพันธ์ ทศทิศรังสรรค์","QA / Tester","DB/discount/injection · program integration"),
("ณฐภาพ สายหล้า","Tech Lead / Architect","Reuse domain · architecture · lint"),
],
"goal":"ปิดงานออกแบบที่ค้างจาก Sprint 1 และยกระดับระบบจาก CLI เป็นโปรแกรม Desktop โดยรักษาฟีเจอร์เดิมทั้งหมด",
"goal_bullets":["ส่วน A: SQLite + Member 4 tiers + Checkout (17 SP)","ส่วน B: GTK4 Desktop Program 4 ส่วน (29 SP)","Migration ต้อง verify จำนวนและมูลค่า 100%","ไม่ให้ regression: v2.0 test suite ต้องผ่านครบ 25 เคส"],
"timeline_title":"สอง Track ใน Sprint 5",
"timeline_subtitle":"พัฒนาฐานโดเมนให้แข็งแรงก่อนต่อด้วย Desktop Presentation Layer",
"timeline":[
("A1","Data Layer","SQLite Singleton · Repository · Migration"),
("A2","Domain","Member 4 tiers · Discount · Checkout"),
("B1","Desktop","GTK4 · Dashboard · Products · Members"),
("B2","Harden","Checkout UI · Tests · Lint · Release Candidate"),
],
"slide4":{"type":"architecture","title":"สถาปัตยกรรม v3.0 + Desktop UI","subtitle":"Presentation layer ใช้ domain เดิมและไม่เก็บ SQL",
"lanes":[
{"title":"CLI / Desktop","body":"Presentation\\nไม่เก็บ SQL","accent":BLUE,"fill":PALE_BLUE},
{"title":"Checkout Service","body":"subtotal\\ndiscount\\ngrand total","accent":CYAN,"fill":PALE_CYAN},
{"title":"Repository + Strategy","body":"InventoryRepository\\nMemberTier × 4","accent":PURPLE,"fill":PALE_PURPLE},
{"title":"SQLite","body":"Products + Members\\ncommit / rollback","accent":GREEN,"fill":PALE_GREEN},
],
"notes":["ทุก query ใช้ parameterized placeholder · products มี CHECK กันค่าติดลบ","program.py reuse InventoryRepository / MemberManager / CheckoutService โดยไม่คัดลอก business logic"]},
"slide5":{"type":"cards","title":"Member Tier + Migration + Checkout","subtitle":"ฟีเจอร์ที่ปิดงานออกแบบค้างจาก Sprint 1",
"cards":[
{"title":"Migration","body":"JSON → SQLite อัตโนมัติเมื่อ DB ว่าง · รองรับคีย์เก่า n/q/p/c · match count/value 100%","accent":CYAN,"fill":PALE_CYAN},
{"title":"4 Member Tiers","body":"Regular 0% · Silver 5% · Gold 10% · Platinum 15% ด้วย Strategy","accent":PURPLE,"fill":PALE_PURPLE},
{"title":"Desktop UI","body":"Dashboard · Products · Members · Checkout รันด้วย python3 program.py","accent":ORANGE,"fill":PALE_ORANGE},
{"title":"Checkout Receipt","body":"subtotal · tier/rate · discount · grand total · remaining stock · low-stock flag","accent":GREEN,"fill":PALE_GREEN},
]},
"slide6":{"type":"metrics","title":"Desktop Program + Quality Gate","subtitle":"GTK4 UI reuse domain — ไม่มีการคัดลอก business logic",
"metrics":[("21/21","Main Tests",GREEN),("25/25","v2 Regression",BLUE),("0","Flake8",CYAN),("0","Bandit",PURPLE)],
"notes":["test_app.py 14 + test_program.py 7 = 21 passed","โปรแกรมเปิดจริงบน GNOME/Wayland","SQLite / Repository / Strategy / CheckoutService ชุดเดิม","UI ไม่ต้องเปิด browser หรือรัน server"]},
"slide7":{"title":"Desktop Program Hardening","subtitle":"ยืนยันทั้ง runtime และ integration ก่อนส่งมอบ",
"cards":[
{"title":"Runtime","body":"python3 program.py เปิด GTK4 window จริงบน GNOME/Wayland","accent":GREEN,"fill":PALE_GREEN},
{"title":"Program Tests","body":"7 integration tests: responsive breakpoints, seed, metrics, LOW/OK, checkout, receipt, CSV","accent":BLUE,"fill":PALE_BLUE},
{"title":"Compatibility","body":"CLI เดิมยังอยู่ใน app.py และ regression v2.0 ผ่าน 25/25","accent":CYAN,"fill":PALE_CYAN},
{"title":"Responsive","body":"Desktop ≥980 · Compact 760–979 · Narrow <760; form/card และ Checkout ปรับตามพื้นที่","accent":ORANGE,"fill":PALE_ORANGE},
]},
"done":["v3.0 SQLite + Member + Checkout","Desktop Program GTK4 4 ส่วน","21 tests + 25 regression · lint/security clean"],
"risks":["ไม่มี Login / Online Payment ตาม scope","SQLite ยังเหมาะกับ local single-user","Desktop build ยังไม่ได้สร้าง release tag เพราะ working tree ยังไม่ commit"],
"nexts":["เพิ่ม tier ใหม่ได้ผ่าน Strategy","ทำ installer/executable ได้ภายหลัง","ทำคู่มือผู้ใช้และปิด UAT sign-off"],
"source":"Source: Phase5/Sprint5/sprint5.md",
"out":"Phase5/Sprint5/sprint5-presentation.pptx"
},
]

if __name__ == "__main__":
    for s in SPRINTS:
        build_sprint(s)
