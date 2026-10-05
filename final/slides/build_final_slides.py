#!/usr/bin/env python3
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageOps
import html

ROOT = Path("/home/panuwat/work/software-project-management")
OUT = ROOT / "final/slides"
IMG = OUT / "img"
OLD = OUT / "img-ระบบเดิม"
JIRA = OUT / "img-jira"
CACHE = OUT / ".cache"
OUT.mkdir(parents=True, exist_ok=True)
CACHE.mkdir(exist_ok=True)

FONT = "Noto Sans"
NAVY="102A3A"; BLUE="0A6D9F"; CYAN="0891B2"; GREEN="237149"
ORANGE="B65319"; RED="A23333"; PURPLE="6D4AFF"; INK="182B36"
MUTED="667B86"; LINE="D9E3E8"; PAPER="F3F6F8"; WHITE="FFFFFF"
PALE_BLUE="EAF4F9"; PALE_GREEN="EEF8F2"; PALE_ORANGE="FFF3E8"
PALE_RED="FFF0F0"; PALE_PURPLE="F2EEFF"

def rgb(h):
    return RGBColor.from_string(h)

def set_bg(slide, color=PAPER):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = rgb(color)

def rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    return shape

def text(slide, value, x, y, w, h, size=18, bold=False, color=INK,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box

def bullets(slide, items, x, y, w, h, size=14, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(7)
        p.line_spacing = 1.08
    return box

def header(slide, kicker, title, subtitle="", accent=BLUE):
    text(slide, kicker.upper(), 0.62, 0.27, 3.3, 0.28, 10.5, True, accent)
    text(slide, title, 0.62, 0.60, 12.0, 0.58, 25, True, NAVY)
    if subtitle:
        text(slide, subtitle, 0.62, 1.13, 12.0, 0.34, 12, False, MUTED)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.62), Inches(12.05), Inches(0.01)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(LINE)
    line.line.fill.background()

def footer(slide, num, source):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(7.04), Inches(12.05), Inches(0.01)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(LINE)
    line.line.fill.background()
    text(slide, source, 0.62, 7.07, 10.9, 0.22, 8.2, False, MUTED)
    text(slide, f"{num:02d}", 11.95, 7.06, 0.70, 0.22, 9, True, BLUE, PP_ALIGN.RIGHT)

def metric(slide, x, y, w, value, label, accent=BLUE):
    rect(slide, x, y, w, 1.08, WHITE, LINE)
    rect(slide, x, y, 0.07, 1.08, accent, accent, False)
    text(slide, value, x+0.20, y+0.10, w-0.30, 0.43, 23, True, accent)
    text(slide, label, x+0.20, y+0.62, w-0.30, 0.26, 10, False, MUTED)

def card(slide, x, y, w, h, title, body, accent=BLUE, fill=WHITE):
    rect(slide, x, y, w, h, fill, LINE)
    rect(slide, x, y, 0.07, h, accent, accent, False)
    text(slide, title, x+0.20, y+0.12, w-0.32, 0.34, 14.5, True, NAVY)
    text(slide, body, x+0.20, y+0.50, w-0.32, h-0.60, 10.5, False, MUTED, valign=MSO_ANCHOR.TOP)

def fit_image(path, width_px, height_px):
    path = Path(path)
    out = CACHE / f"{path.stem}-{width_px}x{height_px}.jpg"
    with Image.open(path).convert("RGB") as im:
        ImageOps.fit(im, (width_px, height_px), method=Image.Resampling.LANCZOS).save(out, quality=92)
    return out

def picture(slide, path, x, y, w, h, crop=False):
    path = Path(path)
    src = fit_image(path, max(500, int(w*140)), max(300, int(h*140))) if crop else path
    slide.shapes.add_picture(str(src), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    frame = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    frame.fill.background()
    frame.line.color.rgb = rgb(LINE)
    frame.line.width = Pt(1)

def new_slide(prs, color=PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, color)
    return slide

def slide_cover(prs):
    s = new_slide(prs, NAVY)
    rect(s, 0, 0, 13.333, 0.11, BLUE, BLUE, False)
    text(s, "SOFTWARE PROJECT MANAGEMENT", 0.72, 0.62, 4.2, 0.30, 10.5, True, "7DD3FC")
    text(s, "วิวัฒนาการระบบจัดการสินค้าคงคลัง", 0.72, 1.18, 6.70, 0.72, 29, True, WHITE)
    text(s, "จาก CLI รุ่นแรก → SQLite → โปรแกรม Desktop UI", 0.72, 1.98, 6.60, 0.44, 17, False, "C8D7E2")
    text(s, "Sprint 1–5 · ตั้งแต่เริ่มโครงการจนถึงตัวส่งมอบสุดท้าย", 0.72, 2.52, 6.45, 0.34, 12, False, "9FB4C0")
    picture(s, IMG/"das.png", 7.66, 0.80, 5.02, 4.92)
    for i, (value, label) in enumerate([
        ("5", "Sprints ปิดครบ"),
        ("95 SP", "Story Points รวม"),
        ("21 + 25", "Current + Regression"),
    ]):
        x = 0.74 + i*2.28
        rect(s, x, 4.45, 2.08, 1.10, "173849", "2C5263")
        text(s, value, x+0.16, 4.59, 1.76, 0.38, 22, True, WHITE)
        text(s, label, x+0.16, 5.03, 1.76, 0.25, 9.2, False, "B8C8D1")
    text(s, "ภานุวัฒน์ ต๋าคำ · เอกพันธ์ ทศทิศรังสรรค์ · ณฐภาพ สายหล้า", 0.74, 6.42, 7.5, 0.30, 10.5, False, "B8C8D1")
    text(s, "ข้อมูลสมมติเพื่อการเรียน", 8.45, 6.42, 4.15, 0.30, 10, False, "8FA6B3", PP_ALIGN.RIGHT)

def slide_team(prs, n):
    s = new_slide(prs)
    header(s, "Project Overview", "ทีม เป้าหมาย และขอบเขต", "1 Sprint ต่อ 1 Phase · พัฒนาต่อเนื่องจากระบบเดิมจนเป็น Desktop Program")
    teams = [
        ("ภานุวัฒน์ ต๋าคำ", "Project Manager / Developer", "บริหารภาพรวม · Data · Domain · UI"),
        ("เอกพันธ์ ทศทิศรังสรรค์", "QA / Tester", "Test Cases · Regression · Edge Cases"),
        ("ณฐภาพ สายหล้า", "Tech Lead / Architect", "Architecture · Code Review · Quality Gate"),
    ]
    y = 1.96
    for name, role, task in teams:
        rect(s, 0.72, y, 5.52, 1.06, WHITE, LINE)
        text(s, name, 0.94, y+0.10, 2.55, 0.30, 14, True, NAVY)
        text(s, role, 0.94, y+0.43, 2.55, 0.22, 10, True, BLUE)
        text(s, task, 3.48, y+0.15, 2.40, 0.48, 10.2, False, MUTED)
        y += 1.22
    card(s, 6.57, 1.96, 5.88, 1.08, "เสถียรภาพ", "Validation · กันค่าติดลบ · ป้องกันข้อมูลเสีย", GREEN, PALE_GREEN)
    card(s, 6.57, 3.21, 5.88, 1.08, "สถาปัตยกรรม", "OOP · Repository · Strategy · แยก UI/Logic", BLUE, PALE_BLUE)
    card(s, 6.57, 4.46, 5.88, 1.08, "วิวัฒนาการ", "JSON → SQLite · Member tiers · Checkout · Desktop Program", PURPLE, PALE_PURPLE)
    footer(s, n, "Source: Sprint 1–5 reports")

def slide_roadmap(prs, n):
    s = new_slide(prs)
    header(s, "Roadmap", "เส้นทาง Sprint 1 → Sprint 5", "ทุก Sprint ปิดแล้ว และแต่ละรอบมีหลักฐานตรวจสอบได้")
    items = [
        ("S1", "W1–W4", "วางฐาน", "15 SP", BLUE),
        ("S2", "W5–W7", "ออกแบบ + Cost", "13 SP", CYAN),
        ("S3", "W8–W11", "Change Control", "16 SP", ORANGE),
        ("S4", "W12", "UAT + Release", "5 SP", GREEN),
        ("S5", "Evolution", "SQLite + Desktop", "46 SP", PURPLE),
    ]
    x0, x1, y = 1.08, 12.15, 3.12
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0), Inches(y), Inches(x1-x0), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = rgb("B9C9D1")
    line.line.fill.background()
    for i, (tag, period, title, sp, accent) in enumerate(items):
        x = x0 + (x1-x0) * (i/(len(items)-1))
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-0.15), Inches(y-0.14), Inches(0.30), Inches(0.30))
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(accent)
        dot.line.color.rgb = rgb(accent)
        text(s, tag, x-0.55, 2.14, 1.10, 0.30, 12.5, True, accent, PP_ALIGN.CENTER)
        text(s, period, x-0.70, 2.52, 1.40, 0.24, 10, False, MUTED, PP_ALIGN.CENTER)
        text(s, title, x-1.0, 3.48, 2.0, 0.38, 12.5, True, NAVY, PP_ALIGN.CENTER)
        text(s, sp, x-0.58, 3.92, 1.16, 0.26, 10.8, True, accent, PP_ALIGN.CENTER)
    metric(s, 2.05, 5.12, 2.62, "95 SP", "Story Points รวม", BLUE)
    metric(s, 5.36, 5.12, 2.62, "5 / 5", "Sprint Closed", GREEN)
    metric(s, 8.67, 5.12, 2.62, "12+", "สัปดาห์ + Evolution", PURPLE)
    footer(s, n, "Source: Sprint 1–5 reports")

def slide_s1_problem(prs, n):
    s = new_slide(prs)
    header(s, "Sprint 1 · Phase 1", "จุดเริ่มต้น: ระบบเดิมแบบ Monolithic", "CLI 5 เมนู · JSON เขียนทับตรง · global state · validation ยังไม่ครบ", BLUE)
    picture(s, OLD/"CLI.png", 0.72, 1.92, 5.55, 3.55)
    card(s, 6.58, 1.92, 5.75, 1.00, "ปัญหาโครงสร้าง", "main() ทำทั้ง UI + Logic + Persistence ในจุดเดียว", RED, PALE_RED)
    card(s, 6.58, 3.08, 5.75, 1.00, "ความเสี่ยงข้อมูล", "เขียน data.json ทับตรงและใช้ global x", ORANGE, PALE_ORANGE)
    card(s, 6.58, 4.24, 5.75, 1.00, "Baseline คุณภาพ", "Pylint 7.10/10 · Cyclomatic Complexity main = 14 (C)", BLUE, PALE_BLUE)
    text(s, "Sprint Goal", 0.76, 5.78, 2.0, 0.30, 14.5, True, NAVY)
    bullets(s, ["Project Charter / Scope", "วิเคราะห์ Hotspot + Blueprint", "Refactor OOP + Automated Tests"], 0.76, 6.06, 11.4, 0.62, 11.8)
    footer(s, n, "Source: Phase1/Sprint1/sprint1.md")

def slide_s1_refactor(prs, n):
    s = new_slide(prs)
    header(s, "Sprint 1 · Phase 1", "Hotspot → Refactor → Automated Test", "OOP 3 คลาส · Validation · Atomic Save · Backward Compatibility", BLUE)
    picture(s, OLD/"Hotspot Flow Diagram.png", 0.72, 1.90, 6.05, 2.48)
    picture(s, OLD/"test_app.png", 0.72, 4.66, 6.05, 1.82)
    card(s, 7.08, 1.92, 5.20, 0.98, "Product", "เก็บ state และรองรับ key เก่า n/q/p/c", BLUE, PALE_BLUE)
    card(s, 7.08, 3.05, 5.20, 0.98, "InventoryManager", "Business Logic + validation + atomic write", CYAN, PALE_BLUE)
    card(s, 7.08, 4.18, 5.20, 0.98, "InventoryCLI", "Presentation layer แยกจาก logic", PURPLE, PALE_PURPLE)
    metric(s, 7.08, 5.46, 2.36, "5 / 5", "PyTest PASS", GREEN)
    metric(s, 9.68, 5.46, 2.60, "15 SP", "SPM-6…10 Done", BLUE)
    footer(s, n, "Source: Phase1/Sprint1/sprint1.md")

def slide_s2(prs, n):
    s = new_slide(prs)
    header(s, "Sprint 2 · Phase 2", "To-Be Architecture + Quality Gate", "ออกแบบก่อนลงมือ เพื่อให้ persistence และ policy เปลี่ยนได้โดยไม่กระทบ UI", CYAN)
    layers = [
        ("Presentation", "CLI / Desktop", "ไม่ถือ SQL", BLUE, PALE_BLUE),
        ("Service", "Checkout / Validation", "Business rules", CYAN, "EAF8FB"),
        ("Repository", "InventoryRepository", "ซ่อน CRUD / SQL", PURPLE, PALE_PURPLE),
        ("Database", "SQLite", "Transaction", GREEN, PALE_GREEN),
    ]
    x = 0.77
    for i, (ttl, body, sub, accent, fill) in enumerate(layers):
        rect(s, x, 2.12, 2.66, 2.22, fill, LINE)
        text(s, ttl, x+0.16, 2.34, 2.34, 0.32, 14, True, accent, PP_ALIGN.CENTER)
        text(s, body, x+0.16, 2.93, 2.34, 0.36, 11.7, True, NAVY, PP_ALIGN.CENTER)
        text(s, sub, x+0.16, 3.46, 2.34, 0.28, 10, False, MUTED, PP_ALIGN.CENTER)
        if i < 3:
            text(s, "→", x+2.70, 2.94, 0.35, 0.34, 20, True, MUTED, PP_ALIGN.CENTER)
        x += 3.0
    metric(s, 0.80, 5.06, 2.68, "18,823 ฿", "Cost Baseline", CYAN)
    metric(s, 3.80, 5.06, 2.68, "36 h", "Capacity / Sprint", BLUE)
    metric(s, 6.80, 5.06, 2.68, "≥85%", "Coverage Target", GREEN)
    metric(s, 9.80, 5.06, 2.68, "≤8", "CC Target", PURPLE)
    footer(s, n, "Source: Phase2/Sprint2/sprint2.md")

def slide_s3(prs, n):
    s = new_slide(prs)
    header(s, "Sprint 3 · Phase 3", "Change Control ที่ตรวจสอบย้อนหลังได้", "CR / Bug ทุกใบเชื่อม Scope, Cost, Architecture และ Test", ORANGE)
    card(s, 0.75, 1.98, 3.75, 1.55, "CR-01", "Barcode + Reorder Point\nImpact +8 man-hours\nส่งใน v2.0", ORANGE, PALE_ORANGE)
    card(s, 4.78, 1.98, 3.75, 1.55, "CR-02", "CSV Export Fast-Track\n4.5 h = 1,350 บาท\nเบิกจาก Reserve", CYAN, "EAF8FB")
    card(s, 8.81, 1.98, 3.75, 1.55, "BUG-101", "KeyError: barcode\nfallback + regression test", RED, PALE_RED)
    text(s, "กระบวนการ 4 ขั้น", 0.82, 3.98, 2.4, 0.32, 16.5, True, NAVY)
    steps = [("1", "Submit"), ("2", "Impact"), ("3", "CCB Review"), ("4", "Baseline")]
    x = 0.82
    for tag, label in steps:
        rect(s, x, 4.48, 2.48, 0.90, WHITE, LINE)
        text(s, tag, x+0.10, 4.62, 0.38, 0.28, 13, True, ORANGE, PP_ALIGN.CENTER)
        text(s, label, x+0.48, 4.61, 1.75, 0.30, 12.5, True, NAVY)
        x += 2.78
    metric(s, 0.82, 5.84, 2.43, "16 SP", "Done", ORANGE)
    metric(s, 3.50, 5.84, 2.43, "27 / 29", "Before Freeze", GREEN)
    metric(s, 6.18, 5.84, 2.43, "WIP 3", "Review Limit", PURPLE)
    metric(s, 8.86, 5.84, 2.43, "1,350 ฿", "Reserve Used", CYAN)
    footer(s, n, "Source: Phase3/Sprint3/sprint3.md")

def slide_s4(prs, n):
    s = new_slide(prs)
    header(s, "Sprint 4 · Phase 4", "UAT และ Release v2.0", "Acceptance scenario + regression + release gate", GREEN)
    metric(s, 0.76, 1.94, 2.64, "25 / 25", "Regression PASS", GREEN)
    metric(s, 3.64, 1.94, 2.64, "3 / 3", "UAT Scenarios", BLUE)
    metric(s, 6.52, 1.94, 2.64, "1.00", "SPI", CYAN)
    metric(s, 9.40, 1.94, 2.64, "−575 ฿", "Cost Variance", ORANGE)
    card(s, 0.76, 3.35, 3.60, 1.48, "SC01 Inventory", "เพิ่มสินค้า 10 · ROP 5 → PASS", GREEN, PALE_GREEN)
    card(s, 4.58, 3.35, 3.60, 1.48, "SC02 Cashier", "ตัด 6 เหลือ 4 → Low-stock → PASS", GREEN, PALE_GREEN)
    card(s, 8.40, 3.35, 3.60, 1.48, "SC03 Purchasing", "Export CSV → PASS", GREEN, PALE_GREEN)
    rect(s, 0.76, 5.27, 11.24, 0.88, WHITE, LINE)
    text(s, "v2.0.0-evolution", 1.00, 5.45, 2.15, 0.30, 13, True, GREEN)
    text(s, "CHANGELOG · Final EVM · Technical Gate ผ่าน", 3.12, 5.44, 6.80, 0.32, 12.4, True, NAVY)
    text(s, "Formal sign-off ยังรอลงนาม", 9.75, 6.27, 2.25, 0.22, 8.8, False, MUTED, PP_ALIGN.RIGHT)
    footer(s, n, "Source: Phase4/Sprint4/sprint4.md")

def slide_evolution(prs, n):
    s = new_slide(prs)
    header(s, "Evolution", "จาก v1.0 ถึง v4.0 Desktop", "แก้ปัญหาทีละชั้นและรักษา backward compatibility")
    versions = [
        ("v1.0", "JSON", "Monolithic", "CLI 5 เมนู", RED, PALE_RED),
        ("v2.0", "JSON + Atomic", "OOP 3 classes", "Validation · Barcode · ROP · CSV", ORANGE, PALE_ORANGE),
        ("v3.0", "SQLite", "Repository + Strategy", "Member 4 tiers · Checkout", CYAN, "EAF8FB"),
        ("v4.0 Desktop", "SQLite", "GTK4 + reuse domain", "Dashboard · Products · Members · Checkout", GREEN, PALE_GREEN),
    ]
    y = 1.94
    for version, db, arch, features, accent, fill in versions:
        rect(s, 0.78, y, 11.80, 1.04, fill, LINE)
        text(s, version, 1.02, y+0.15, 1.60, 0.34, 15.2, True, accent)
        text(s, db, 2.72, y+0.15, 1.85, 0.34, 11.8, True, NAVY)
        text(s, arch, 4.64, y+0.15, 2.95, 0.34, 11.6, True, NAVY)
        text(s, features, 7.80, y+0.13, 4.34, 0.50, 10.4, False, MUTED)
        y += 1.18
    footer(s, n, "Source: presentation.md")

def slide_s5_arch(prs, n):
    s = new_slide(prs)
    header(s, "Sprint 5 · Phase 5", "Data/Domain Architecture", "Desktop UI reuse domain เดิม ไม่เขียน SQL หรือคำนวณส่วนลดซ้ำ", PURPLE)
    layers = [
        ("CLI / Desktop", "Presentation", "app.py / program.py", BLUE, PALE_BLUE),
        ("CheckoutService", "Service", "subtotal · discount · total", CYAN, "EAF8FB"),
        ("Repository + Strategy", "Domain", "InventoryRepository · MemberTier", PURPLE, PALE_PURPLE),
        ("SQLite", "Persistence", "products · members · transaction", GREEN, PALE_GREEN),
    ]
    x = 0.70
    for i, (title, level, detail, accent, fill) in enumerate(layers):
        rect(s, x, 2.06, 2.70, 2.30, fill, LINE)
        text(s, title, x+0.16, 2.31, 2.38, 0.34, 13.6, True, accent, PP_ALIGN.CENTER)
        text(s, level, x+0.16, 2.92, 2.38, 0.30, 11.2, True, NAVY, PP_ALIGN.CENTER)
        text(s, detail, x+0.16, 3.42, 2.38, 0.56, 9.9, False, MUTED, PP_ALIGN.CENTER)
        if i < 3:
            text(s, "→", x+2.72, 3.00, 0.35, 0.34, 20, True, MUTED, PP_ALIGN.CENTER)
        x += 3.03
    metric(s, 0.80, 5.05, 2.68, "17 SP", "Data + Domain", PURPLE)
    metric(s, 3.80, 5.05, 2.68, "29 SP", "Desktop UI", BLUE)
    metric(s, 6.80, 5.05, 2.68, "46 SP", "Sprint 5 Total", GREEN)
    metric(s, 9.80, 5.05, 2.68, "100%", "Migration Verify", CYAN)
    footer(s, n, "Source: Phase5/Sprint5/sprint5.md")

def slide_member_checkout(prs, n):
    s = new_slide(prs)
    header(s, "Sprint 5 · Phase 5", "Member Strategy + Checkout", "CheckoutService เป็น source of truth เดียว", PURPLE)
    tiers = [
        ("Regular", "0%", BLUE),
        ("Silver", "5%", CYAN),
        ("Gold", "10%", ORANGE),
        ("Platinum", "15%", PURPLE),
    ]
    for i, (name, rate, accent) in enumerate(tiers):
        x = 0.78 + i*3.02
        rect(s, x, 1.96, 2.70, 1.40, WHITE, LINE)
        text(s, name, x+0.18, 2.14, 2.34, 0.32, 13.8, True, NAVY, PP_ALIGN.CENTER)
        text(s, rate, x+0.18, 2.58, 2.34, 0.42, 23, True, accent, PP_ALIGN.CENTER)
    rect(s, 0.78, 3.78, 5.50, 2.04, PALE_BLUE, LINE)
    text(s, "ตัวอย่าง Gold", 1.04, 4.00, 4.8, 0.34, 16.5, True, NAVY)
    bullets(s, ["2 × 500 = 1,000 บาท", "ส่วนลด 10% = 100 บาท", "ยอดสุทธิ = 900 บาท", "Guest = ราคาเต็ม"], 1.05, 4.42, 4.85, 1.22, 12.6)
    rect(s, 6.65, 3.78, 5.70, 2.04, PALE_GREEN, LINE)
    text(s, "ผลต่อการออกแบบ", 6.93, 4.00, 5.05, 0.34, 16.5, True, GREEN)
    bullets(s, ["เพิ่ม tier ใหม่โดยไม่กระจาย if-else", "UI ไม่คำนวณ discount ซ้ำ", "CheckoutService เป็น truth เดียว", "Receipt แสดง stock คงเหลือ"], 6.95, 4.42, 5.02, 1.22, 12.3)
    footer(s, n, "Source: Phase5/Sprint5/sprint5.md")

def slide_jira(prs, n):
    s = new_slide(prs)
    header(s, "Jira · Sprint 5", "Epic SPM-28: CLI → Desktop Program UI", "คง Story Points / Parent / Done เดิม แต่แก้ scope ให้ตรงกับตัวส่งมอบจริง", PURPLE)
    picture(s, JIRA/"image.png", 0.72, 1.88, 7.10, 4.55)
    card(s, 8.08, 1.90, 4.40, 1.04, "SPM-29 · Foundation", "GTK4 scaffold + reuse domain", PURPLE, PALE_PURPLE)
    card(s, 8.08, 3.08, 4.40, 1.04, "SPM-30 · Products", "Dashboard + Products Desktop UI", BLUE, PALE_BLUE)
    card(s, 8.08, 4.26, 4.40, 1.04, "SPM-31 · Checkout", "Members + Checkout Desktop UI", CYAN, "EAF8FB")
    card(s, 8.08, 5.44, 4.40, 1.04, "SPM-32 · Hardening", "Responsive + Quality Gate", GREEN, PALE_GREEN)
    footer(s, n, "Source: Jira SPM-28…32")

def slide_ui(prs, n, title, subtitle, image_name, items, accent):
    s = new_slide(prs)
    header(s, "Final Desktop Program", title, subtitle, accent)
    picture(s, IMG/image_name, 0.68, 1.88, 8.25, 4.65)
    rect(s, 9.23, 1.88, 3.45, 4.65, WHITE, LINE)
    text(s, "จุดสำคัญ", 9.50, 2.17, 2.85, 0.34, 16.5, True, NAVY)
    bullets(s, items, 9.50, 2.70, 2.78, 3.15, 12.4)
    footer(s, n, "Source: final/desktop-app · Screenshots")

def slide_quality(prs, n):
    s = new_slide(prs)
    header(s, "Quality Gate", "สถานะสุดท้ายของโครงการ", "Domain + Desktop integration + regression + static checks", GREEN)
    metric(s, 0.72, 1.94, 2.72, "21 / 21", "Main Tests", GREEN)
    metric(s, 3.72, 1.94, 2.72, "25 / 25", "v2 Regression", BLUE)
    metric(s, 6.72, 1.94, 2.72, "0", "Flake8 Issues", CYAN)
    metric(s, 9.72, 1.94, 2.72, "0", "Bandit Issues", PURPLE)
    rect(s, 0.72, 3.42, 5.82, 2.48, WHITE, LINE)
    text(s, "Desktop Integration 7 เคส", 0.98, 3.67, 5.20, 0.34, 16.5, True, NAVY)
    bullets(s, ["Responsive breakpoints", "Seed + Dashboard metrics", "LOW / OK status", "Member checkout + Receipt", "CSV export"], 1.00, 4.08, 5.10, 1.45, 12.1)
    rect(s, 6.80, 3.42, 5.72, 2.48, WHITE, LINE)
    text(s, "Regression v2.0", 7.07, 3.67, 5.10, 0.34, 16.5, True, NAVY)
    bullets(s, ["Baseline behavior", "CR-01 Barcode / ROP", "CR-02 CSV Export", "BUG-101 legacy JSON", "CLI + Integration"], 7.08, 4.08, 5.02, 1.45, 12.1)
    footer(s, n, "Source: test_app.py · test_program.py")

def slide_pm(prs, n):
    s = new_slide(prs)
    header(s, "Project Management", "มุมมองรวม 5 Sprint", "Scope · Cost · Quality · Change · Release เชื่อมด้วยหลักฐานเดียวกัน", ORANGE)
    data = [
        ("Scope", "95 SP · ทุก Sprint ปิดแล้ว", BLUE, PALE_BLUE),
        ("Cost", "Baseline 18,823 ฿ · Final CV −575 ฿", ORANGE, PALE_ORANGE),
        ("Quality", "DoR/DoD · ISO Gate · Automated Tests", GREEN, PALE_GREEN),
        ("Change", "CR-01 · CR-02 · BUG-101", PURPLE, PALE_PURPLE),
        ("Release", "v2.0 gate + Desktop candidate", CYAN, "EAF8FB"),
        ("Evidence", "Reports · Jira · Tests · Screenshots", BLUE, WHITE),
    ]
    for i, (title, body, accent, fill) in enumerate(data):
        col = i % 3
        row = i // 3
        card(s, 0.78+col*4.05, 1.98+row*2.12, 3.65, 1.68, title, body, accent, fill)
    footer(s, n, "Source: Sprint 1–5 reports")

def slide_close(prs, n):
    s = new_slide(prs, NAVY)
    rect(s, 0, 0, 13.333, 0.11, BLUE, BLUE, False)
    text(s, "PROJECT OUTCOME", 0.75, 0.70, 3.0, 0.28, 10.5, True, "7DD3FC")
    text(s, "จาก CLI ที่เปราะบาง\nสู่โปรแกรม Desktop ที่ทดสอบได้", 0.75, 1.30, 7.20, 1.35, 29, True, WHITE, valign=MSO_ANCHOR.TOP)
    text(s, "สิ่งที่พิสูจน์ได้จากหลักฐาน", 0.78, 3.08, 3.2, 0.34, 16, True, "C8D7E2")
    bullets(s, ["Monolithic → Layered Architecture", "JSON → SQLite + Transaction", "Member Strategy + Checkout", "Desktop UI reuse domain", "21/21 + 25/25 · Flake8 0 · Bandit 0"], 0.80, 3.48, 6.55, 2.10, 12.8, "E4EDF2")
    rect(s, 8.10, 1.52, 4.50, 3.72, "173849", "2C5263")
    text(s, "Next Step", 8.40, 1.84, 3.8, 0.36, 18, True, "7DD3FC")
    bullets(s, ["Executable / Installer", "Formal UAT Sign-off", "Role/Login หากขยาย scope", "PostgreSQL เมื่อ multi-user"], 8.40, 2.38, 3.65, 2.0, 13.4, WHITE)
    text(s, "ขอบคุณครับ", 8.40, 5.78, 3.80, 0.42, 23, True, WHITE)
    text(s, "Software Project Management · Final Presentation", 0.75, 6.62, 5.5, 0.26, 10, False, "8FA6B3")
    text(s, f"{n:02d}", 11.90, 6.62, 0.70, 0.26, 10, True, "7DD3FC", PP_ALIGN.RIGHT)

def build_html():
    slides = [
        ("Software Project Management", "วิวัฒนาการระบบจัดการสินค้าคงคลัง", "จาก CLI รุ่นแรก → SQLite → โปรแกรม Desktop UI", "img/das.png"),
        ("Project Overview", "ทีม เป้าหมาย และขอบเขต", "ภานุวัฒน์ ต๋าคำ · เอกพันธ์ ทศทิศรังสรรค์ · ณฐภาพ สายหล้า", None),
        ("Roadmap", "Sprint 1 → Sprint 5", "15 SP → 13 SP → 16 SP → 5 SP → 46 SP · รวม 95 SP", None),
        ("Sprint 1", "ระบบเดิมแบบ Monolithic", "CLI 5 เมนู · JSON เขียนทับตรง · global state · CC main = 14", "img-ระบบเดิม/CLI.png"),
        ("Sprint 1", "Hotspot → Refactor → Automated Test", "OOP 3 คลาส · Atomic Save · Validation · PyTest 5/5", "img-ระบบเดิม/Hotspot Flow Diagram.png"),
        ("Sprint 2", "To-Be Architecture + Quality Gate", "Presentation → Service → Repository → SQLite · Cost Baseline 18,823 บาท", None),
        ("Sprint 3", "Change Control", "CR-01 Barcode/ROP · CR-02 CSV · BUG-101 legacy JSON", None),
        ("Sprint 4", "UAT + Release v2.0", "25/25 Regression · UAT 3/3 · SPI 1.00 · CV -575 บาท", None),
        ("Evolution", "v1.0 → v4.0 Desktop", "Monolithic JSON → OOP → SQLite/Strategy → GTK4 Desktop", None),
        ("Sprint 5", "Data/Domain Architecture", "17 SP Data/Domain + 29 SP Desktop UI = 46 SP", None),
        ("Sprint 5", "Member Strategy + Checkout", "Regular 0% · Silver 5% · Gold 10% · Platinum 15%", None),
        ("Jira", "Epic SPM-28", "SPM-29…32 เปลี่ยน scope เป็น Desktop Program UI และยัง Done", "img-jira/image.png"),
        ("Final Desktop Program", "Dashboard", "ภาพรวมคลัง · low-stock · Responsive", "img/das.png"),
        ("Final Desktop Program", "Products", "เพิ่ม/แก้ไข · Search · LOW/OK · CSV", "img/product.png"),
        ("Final Desktop Program", "Members", "CRUD สมาชิก · 4 Tiers · Discount badge", "img/member.png"),
        ("Final Desktop Program", "Checkout", "Guest/Member · Strategy discount · Receipt", "img/checkout.png"),
        ("Quality Gate", "สถานะสุดท้าย", "21/21 Main · 25/25 Regression · Flake8 0 · Bandit 0", None),
        ("Project Management", "มุมมองรวม 5 Sprint", "Scope · Cost · Quality · Change · Release · Evidence", None),
        ("Project Outcome", "จาก CLI ที่เปราะบางสู่ Desktop Program ที่ทดสอบได้", "Next: Installer · Formal UAT · Role/Login · PostgreSQL", None),
    ]
    cards = []
    for i, (kicker, title, body, image) in enumerate(slides, start=1):
        image_html = f'<img src="{html.escape(image)}" alt="{html.escape(title)}">' if image else ""
        cards.append(
            f'''<section class="slide" id="slide-{i}">
              <div class="kicker">{html.escape(kicker)}</div>
              <h1>{html.escape(title)}</h1>
              <p>{html.escape(body)}</p>
              {image_html}
              <div class="foot"><span>Sprint 1–5 · Final Presentation</span><b>{i:02d}</b></div>
            </section>'''
        )
    css = '''
    :root{font-family:"Noto Sans","Noto Sans Thai",sans-serif;color:#182b36;background:#e7eef2}
    *{box-sizing:border-box} body{margin:0}
    .toolbar{width:min(1280px,calc(100% - 32px));margin:12px auto;display:flex;justify-content:space-between;align-items:center;color:#667b86}
    .deck{width:min(1280px,calc(100vw - 32px));aspect-ratio:16/9;margin:auto;background:#fff;box-shadow:0 20px 50px #102a3a22;overflow:hidden}
    .slide{display:none;width:100%;height:100%;padding:5% 5%;position:relative;background:#f3f6f8}
    .slide.active{display:block}.kicker{font-size:14px;font-weight:800;color:#0a6d9f;text-transform:uppercase;letter-spacing:.06em}
    h1{margin:20px 0 12px;font-size:clamp(32px,4vw,60px);line-height:1.08;color:#102a3a;max-width:90%}
    p{font-size:clamp(18px,2vw,28px);line-height:1.45;color:#667b86;max-width:85%;margin:0 0 24px}
    img{display:block;max-width:82%;max-height:56%;object-fit:contain;border:1px solid #d9e3e8;border-radius:14px;background:white}
    .foot{position:absolute;left:5%;right:5%;bottom:3.5%;border-top:1px solid #d9e3e8;padding-top:10px;display:flex;justify-content:space-between;color:#667b86;font-size:12px}
    .controls{width:min(1280px,calc(100% - 32px));margin:12px auto;display:flex;justify-content:space-between}
    button{font:700 14px "Noto Sans";padding:10px 18px;border-radius:8px;border:1px solid #d9e3e8;background:white;color:#102a3a}
    button:disabled{opacity:.35}
    @media print{@page{size:13.333in 7.5in;margin:0}.toolbar,.controls{display:none}.deck{width:13.333in;aspect-ratio:auto;box-shadow:none;overflow:visible}.slide{display:block!important;width:13.333in;height:7.5in;break-after:page}}
    '''
    body = "\n".join(cards)
    doc = f'''<!doctype html><html lang="th"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
    <title>Final Sprint 1–5 Presentation</title><style>{css}</style></head><body>
    <div class="toolbar"><strong>Final Presentation · Sprint 1–5</strong><span id="counter"></span></div>
    <main class="deck">{body}</main>
    <div class="controls"><button id="prev">← ก่อนหน้า</button><button id="next">ถัดไป →</button></div>
    <script>
    const slides=[...document.querySelectorAll(".slide")];let i=0;
    const prev=document.getElementById("prev"),next=document.getElementById("next"),counter=document.getElementById("counter");
    function show(n){{i=Math.max(0,Math.min(n,slides.length-1));slides.forEach((s,j)=>s.classList.toggle("active",i===j));counter.textContent=(i+1)+" / "+slides.length;prev.disabled=i===0;next.disabled=i===slides.length-1;}}
    prev.onclick=()=>show(i-1);next.onclick=()=>show(i+1);
    addEventListener("keydown",e=>{{if(["ArrowRight","PageDown"," "].includes(e.key))show(i+1);if(["ArrowLeft","PageUp"].includes(e.key))show(i-1);}});
    show(0);
    </script></body></html>'''
    (OUT/"slides.html").write_text(doc, encoding="utf-8")

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_cover(prs)
    slide_team(prs, 2)
    slide_roadmap(prs, 3)
    slide_s1_problem(prs, 4)
    slide_s1_refactor(prs, 5)
    slide_s2(prs, 6)
    slide_s3(prs, 7)
    slide_s4(prs, 8)
    slide_evolution(prs, 9)
    slide_s5_arch(prs, 10)
    slide_member_checkout(prs, 11)
    slide_jira(prs, 12)
    slide_ui(prs, 13, "Dashboard", "เห็นสถานะคลังและรายการที่ควรจัดการก่อน", "das.png",
             ["ประเภทสินค้า / มูลค่าคงคลัง / low-stock", "SQLite status ชัดเจน", "Responsive layout", "เวลาอัปเดตล่าสุด"], BLUE)
    slide_ui(prs, 14, "Products", "เพิ่ม แก้ไข ค้นหา ตัดสต็อก ลบ และ Export CSV", "product.png",
             ["Form เพิ่ม/แก้ไข", "Search รหัส / ชื่อ / หมวดหมู่", "LOW / OK badge", "Export CSV + action ต่อแถว"], CYAN)
    slide_ui(prs, 15, "Members", "จัดการสมาชิกและระดับส่วนลด", "member.png",
             ["CRUD สมาชิก", "Regular / Silver / Gold / Platinum", "Discount badge", "CheckoutService ใช้ policy เดียว"], PURPLE)
    slide_ui(prs, 16, "Checkout", "ตัดสต็อก คิดส่วนลด และแสดงใบเสร็จ", "checkout.png",
             ["Guest / Member checkout", "Strategy discount", "Grand Total เด่น", "แสดง stock คงเหลือ"], GREEN)
    slide_quality(prs, 17)
    slide_pm(prs, 18)
    slide_close(prs, 19)
    prs.save(OUT/"final-sprint-1-to-5.pptx")
    build_html()
    (OUT/"README.md").write_text(
        "# Final Slides — Sprint 1 ถึง Sprint 5\n\n"
        "- final-sprint-1-to-5.pptx — PowerPoint รวม 19 หน้า\n"
        "- final-sprint-1-to-5.pdf — PDF จาก PowerPoint\n"
        "- slides.html — HTML deck\n"
        "- slides.pdf — PDF จาก HTML\n"
        "- sprints/ — ชุดสไลด์แยก Sprint 1–5\n"
        "- build_final_slides.py — generator\n\n"
        "ลำดับ: Sprint 1 วางฐาน → Sprint 2 Architecture/Cost → Sprint 3 Change Control → "
        "Sprint 4 UAT/Release → Sprint 5 SQLite/Desktop → Quality Gate → Outcome\n",
        encoding="utf-8"
    )
    print("CREATED", OUT/"final-sprint-1-to-5.pptx", len(prs.slides))
    print("CREATED", OUT/"slides.html")

if __name__ == "__main__":
    main()
